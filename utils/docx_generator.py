import io
from docx import Document
from docx.shared import Inches


def markdown_to_docx(markdown_text: str) -> io.BytesIO:
    """
    Convert a standard Markdown resume into a clean, professionally formatted 
    Microsoft Word Document (.docx) returned in a bytes buffer.
    """
    doc = Document()

    # Set standard professional margins (1 inch / 72pt on all sides)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    lines = markdown_text.split('\n')
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # 1. Headings (Level 1, 2, 3)
        if stripped.startswith('# '):
            text = stripped[2:].replace('**', '')  # Clean bold markers inside headings
            doc.add_heading(text, level=1)
            
        elif stripped.startswith('## '):
            text = stripped[3:].replace('**', '')
            doc.add_heading(text, level=2)
            
        elif stripped.startswith('### '):
            text = stripped[4:].replace('**', '')
            doc.add_heading(text, level=3)

        # 2. Bullet Lists
        elif stripped.startswith('- ') or stripped.startswith('* '):
            # Native list bullet style
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:]
            
            # Inline bold emphasis parsing (e.g. **Skill**: detail)
            parts = text.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 == 1:
                    run.bold = True

        # 3. Standard Paragraphs
        else:
            p = doc.add_paragraph()
            parts = stripped.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 == 1:
                    run.bold = True

    # Save document directly into an in-memory bytes buffer
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)
    return docx_buffer
